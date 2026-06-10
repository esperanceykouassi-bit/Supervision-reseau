# -*- coding: utf-8 -*-
"""
generer_pptx.py
===============
Génère le support de soutenance au format PowerPoint (.pptx) ÉDITABLE.

Le fichier produit (« presentation-projet.pptx ») suit le canevas
mémoire-projet en 12 sections. Chaque diapositive contient un titre, un
contenu structuré (puces / tableaux) et les NOTES DE L'ORATEUR.

Utilisation :
    pip install python-pptx
    python generer_pptx.py
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

# Racine du dépôt (pour retrouver les images des diagrammes).
RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------------- #
# Charte graphique — inspirée du modèle « Soutenance_Supervision_Reseau » #
#   Bleu marine + sarcelle + menthe, fonds clairs, Cambria/Calibri.       #
# ---------------------------------------------------------------------- #
NAVY       = RGBColor(0x0A, 0x23, 0x42)   # marine principal (titres, panneaux)
NAVY2      = RGBColor(0x13, 0x3B, 0x5C)   # marine décoratif
TEAL_DK    = RGBColor(0x1C, 0x7C, 0x84)   # sarcelle (accents, sous-titres)
TEAL       = RGBColor(0x2E, 0xC4, 0xB6)   # sarcelle vif (barres, icônes)
MINT       = RGBColor(0x5F, 0xE3, 0xD3)   # menthe (numéros, kicker, accents)
BG         = RGBColor(0xF4, 0xF8, 0xFB)   # fond clair des diapos
TXT        = RGBColor(0x0F, 0x1B, 0x2D)   # texte courant
MUTED      = RGBColor(0x64, 0x74, 0x8B)   # texte secondaire / pied de page
LIGHTTXT   = RGBColor(0xC7, 0xD7, 0xE2)   # texte clair sur fond marine

# Alias conservés pour ne pas casser le contenu existant
BLEU_FONCE = NAVY        # structures / titres
BLEU       = TEAL_DK     # accents (lisible sur fond clair)
GRIS       = TXT         # texte courant
GRIS_CLAIR = RGBColor(0xE8, 0xEE, 0xF4)   # lignes alternées / fonds doux
BLANC      = RGBColor(0xFF, 0xFF, 0xFF)
ROUGE      = RGBColor(0xD6, 0x28, 0x28)   # danger / sévérité critique
VERT       = RGBColor(0x1C, 0x7C, 0x84)   # succès → harmonisé en sarcelle

FOOTER_TXT = "Automatisation de la Supervision et de la Détection des Pannes Réseau — OPEN MOISE"

# Format 16:9
LARGEUR = Inches(13.333)
HAUTEUR = Inches(7.5)

prs = Presentation()
prs.slide_width = LARGEUR
prs.slide_height = HAUTEUR

LAYOUT_VIDE = prs.slide_layouts[6]  # diapositive vierge


# ---------------------------------------------------------------------- #
# Fonctions utilitaires de mise en page                                  #
# ---------------------------------------------------------------------- #
def ajouter_diapo():
    return prs.slides.add_slide(LAYOUT_VIDE)


def fond(slide, couleur=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = couleur


def rectangle(slide, x, y, w, h, couleur):
    from pptx.enum.shapes import MSO_SHAPE
    forme = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    forme.fill.solid()
    forme.fill.fore_color.rgb = couleur
    forme.line.fill.background()
    forme.shadow.inherit = False
    return forme


def zone_texte(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    return tb


def style_run(run, taille=18, gras=False, couleur=GRIS, italique=False, police="Calibri"):
    run.font.size = Pt(taille)
    run.font.bold = gras
    run.font.italic = italique
    run.font.color.rgb = couleur
    run.font.name = police


def notes(slide, texte):
    slide.notes_slide.notes_text_frame.text = texte


def carre(slide, x, y, c, taille=Inches(0.5)):
    """Petit carré arrondi de couleur `c` (badge / icône)."""
    from pptx.enum.shapes import MSO_SHAPE
    f = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, taille, taille)
    f.fill.solid(); f.fill.fore_color.rgb = c
    f.line.fill.background(); f.shadow.inherit = False
    return f


def cercle(slide, x, y, d, c):
    """Cercle décoratif plein (pour la page de titre / dividers)."""
    from pptx.enum.shapes import MSO_SHAPE
    f = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    f.fill.solid(); f.fill.fore_color.rgb = c
    f.line.fill.background(); f.shadow.inherit = False
    return f


def pied(slide):
    """Pied de page : intitulé à gauche + numéro de diapo à droite."""
    no = len(prs.slides._sldIdLst)  # index 1-based de la diapo courante
    tb = zone_texte(slide, Inches(0.6), Inches(7.12), Inches(10.5), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = FOOTER_TXT
    style_run(r, 9, False, MUTED)
    tn = zone_texte(slide, Inches(12.2), Inches(7.12), Inches(0.6), Inches(0.3))
    p = tn.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(no)
    style_run(r, 10, False, MUTED)


def bandeau_titre(slide, numero, titre, sous_titre=None):
    """En-tête : badge marine numéroté + titre Cambria + accent + pied de page.

    Reproduit le motif du modèle de référence (fond clair, pas de bandeau plein).
    """
    # Badge carré marine avec le numéro (ou « • ») en menthe.
    carre(slide, Inches(0.55), Inches(0.45), NAVY, taille=Inches(0.8))
    bt = zone_texte(slide, Inches(0.55), Inches(0.45), Inches(0.8), Inches(0.8))
    bt.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = bt.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(numero) if numero is not None else "•"
    style_run(r, 28, True, MINT, police="Cambria")
    # Titre (Cambria, marine).
    tb = zone_texte(slide, Inches(1.6), Inches(0.45), Inches(11.1), Inches(0.8))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = titre
    style_run(r, 26, True, NAVY, police="Cambria")
    # Accent sarcelle court sous le titre.
    rectangle(slide, Inches(1.62), Inches(1.28), Inches(1.5), Inches(0.045), TEAL)
    pied(slide)


def puces(slide, items, x=Inches(0.7), y=Inches(1.5),
          w=Inches(11.9), h=Inches(5.4), taille=20, interligne=8):
    """Ajoute une liste à puces. `items` = liste de (texte, niveau, gras, couleur)."""
    tb = zone_texte(slide, x, y, w, h)
    tf = tb.text_frame
    first = True
    for item in items:
        texte = item[0]
        niveau = item[1] if len(item) > 1 else 0
        gras = item[2] if len(item) > 2 else False
        couleur = item[3] if len(item) > 3 else GRIS
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = niveau
        p.space_after = Pt(interligne)
        puce = "•  " if niveau == 0 else "–  "
        if texte == "":
            puce = ""
        r = p.add_run(); r.text = puce + texte
        style_run(r, taille - niveau * 2, gras, couleur)
    return tb


def tableau(slide, donnees, x=Inches(0.7), y=Inches(1.5),
            w=Inches(11.9), h=Inches(5.0), taille=15, surligne_derniere_col=False):
    """Crée un tableau stylé. `donnees` = liste de lignes (listes de cellules)."""
    nb_lignes = len(donnees)
    nb_cols = len(donnees[0])
    g = slide.shapes.add_table(nb_lignes, nb_cols, x, y, w, h).table
    for j, cell_txt in enumerate(donnees[0]):
        cell = g.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = BLEU_FONCE
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(cell_txt)
        style_run(r, taille, True, BLANC)
    for i in range(1, nb_lignes):
        for j in range(nb_cols):
            cell = g.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLANC if i % 2 else GRIS_CLAIR
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(donnees[i][j])
            couleur = GRIS
            gras = False
            if surligne_derniere_col and j == nb_cols - 1:
                couleur = VERT; gras = True
            style_run(r, taille, gras, couleur)
    return g


def encadre(slide, texte, x=Inches(0.7), y=Inches(5.9),
            w=Inches(11.9), h=Inches(1.0), fond_c=None, barre=TEAL):
    """Encadré « message clé » avec barre latérale sarcelle.

    La couleur du texte s'adapte automatiquement à la clarté du fond.
    """
    if fond_c is None:
        fond_c = RGBColor(0xEA, 0xF3, 0xF4)        # tuile très claire (défaut)
    fonds_sombres = {NAVY, NAVY2, TEAL_DK}
    couleur_txt = LIGHTTXT if fond_c in fonds_sombres else NAVY
    rectangle(slide, x, y, Inches(0.12), h, barre)
    boite = rectangle(slide, x + Inches(0.12), y, w - Inches(0.12), h, fond_c)
    tf = boite.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = texte
    style_run(r, 16, True, couleur_txt, italique=True, police="Cambria")


def _dim_ajustee(chemin, max_w, max_h):
    """Retourne (largeur, hauteur) en EMU pour faire tenir l'image en gardant
    le ratio d'aspect dans la boîte (max_w x max_h)."""
    iw, ih = Image.open(chemin).size
    aspect = iw / ih
    w = max_w
    h = int(w / aspect)
    if h > max_h:
        h = max_h
        w = int(h * aspect)
    return w, h


def diapo_image(numero, titre, chemin, legende="", note_txt="",
                max_w=Inches(11.8), max_h=Inches(5.3), top=Inches(1.35),
                fond_clair=True):
    """Crée une diapositive « titre + image centrée » (diagramme pleine page)."""
    s = ajouter_diapo(); fond(s, BG)
    bandeau_titre(s, numero, titre)
    chemin = os.path.join(RACINE, chemin)
    w, h = _dim_ajustee(chemin, int(max_w), int(max_h))
    left = int((int(LARGEUR) - w) / 2)
    s.shapes.add_picture(chemin, left, top, width=w, height=h)
    if legende:
        tb = zone_texte(s, Inches(0.5), Inches(6.72), Inches(11.4), Inches(0.4))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = legende
        style_run(r, 12, False, MUTED, italique=True)
    if note_txt:
        notes(s, note_txt)
    return s


def diapo_deux_images(numero, titre, g_chemin, g_lbl, d_chemin, d_lbl, note_txt=""):
    """Crée une diapositive avec deux images côte à côte (ex. MCD | MLD)."""
    s = ajouter_diapo(); fond(s)
    bandeau_titre(s, numero, titre)
    for chemin, lbl, x in [(g_chemin, g_lbl, Inches(0.4)),
                           (d_chemin, d_lbl, Inches(6.9))]:
        # libellé
        tb = zone_texte(s, x, Inches(1.3), Inches(6.0), Inches(0.4))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = lbl
        style_run(r, 16, True, BLEU_FONCE)
        # image
        ch = os.path.join(RACINE, chemin)
        w, h = _dim_ajustee(ch, int(Inches(5.9)), int(Inches(5.0)))
        left = int(x) + int((int(Inches(6.0)) - w) / 2)
        s.shapes.add_picture(ch, left, Inches(1.75), width=w, height=h)
    if note_txt:
        notes(s, note_txt)
    return s


def diapo_grille_images(numero, titre, items, note_txt=""):
    """Grille d'aperçus (4 colonnes). items = liste de (chemin, label)."""
    s = ajouter_diapo(); fond(s)
    bandeau_titre(s, numero, titre)
    cols = 4
    cell_w = Inches(3.0)
    cell_h = Inches(2.35)
    x0, y0 = Inches(0.35), Inches(1.45)
    gx, gy = Inches(0.18), Inches(0.55)
    for idx, (chemin, label) in enumerate(items):
        col = idx % cols
        row = idx // cols
        cx = int(x0) + col * (int(cell_w) + int(gx))
        cy = int(y0) + row * (int(cell_h) + int(gy))
        ch = os.path.join(RACINE, chemin)
        w, h = _dim_ajustee(ch, int(cell_w), int(cell_h))
        left = cx + int((int(cell_w) - w) / 2)
        s.shapes.add_picture(ch, left, cy, width=w, height=h)
        tb = zone_texte(s, cx, cy + int(cell_h) + Inches(0.02), cell_w, Inches(0.4))
        p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        style_run(r, 11, True, BLEU_FONCE)
    if note_txt:
        notes(s, note_txt)
    return s


# ====================================================================== #
#  DIAPO 0 — TITRE                                                       #
# ====================================================================== #
s = ajouter_diapo(); fond(s, NAVY)
# Cercles décoratifs (coins)
cercle(s, Inches(10.8), Inches(-1.4), Inches(3.8), NAVY2)
cercle(s, Inches(11.9), Inches(-0.5), Inches(1.7), TEAL_DK)
cercle(s, Inches(-1.0), Inches(5.3), Inches(3.4), NAVY2)
cercle(s, Inches(0.8), Inches(1.1), Inches(1.0), TEAL)
# Kicker
tb = zone_texte(s, Inches(0.8), Inches(2.15), Inches(11.6), Inches(0.45))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "MÉMOIRE DE PROJET DE FIN D'ÉTUDES — MASTER 2 RIT"
style_run(r, 15, True, MINT, police="Cambria")
# Titre principal
tb = zone_texte(s, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.9))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Automatisation de la Supervision et de la Détection des Pannes Réseau"
style_run(r, 42, True, BLANC, police="Cambria")
# Barre d'accent
rectangle(s, Inches(0.8), Inches(4.75), Inches(2.6), Inches(0.06), TEAL)
# Sous-titre
tb = zone_texte(s, Inches(0.8), Inches(4.95), Inches(11.5), Inches(0.6))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "Une solution open-source déployée au sein d'OPEN MOISE (ESN d'infogérance)"
style_run(r, 16, False, LIGHTTXT)
# Bloc auteur
tb = zone_texte(s, Inches(0.8), Inches(6.05), Inches(11.6), Inches(1.1))
tf = tb.text_frame
for i, (txt, sz, g, c) in enumerate([
    ("Présenté par : [NOM Prénom de l'étudiant]      Encadré par : [Directeur de mémoire]", 13, True, BLANC),
    ("Structure d'accueil : OPEN MOISE", 12.5, False, LIGHTTXT),
    ("Année académique 2025 – 2026", 12, True, MINT),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = txt
    style_run(r, sz, g, c)
notes(s, "Bonjour Mesdames et Messieurs les membres du jury. Je vous remercie de "
         "votre présence. Je vais vous présenter mon mémoire-projet : une solution "
         "d'automatisation de la supervision et de la détection des pannes réseau, "
         "pensée comme un véritable produit, du besoin du marché au modèle économique.")

# ====================================================================== #
#  DIAPO 1 — PLAN                                                        #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, None, "Plan de la présentation")
gauche = [
    ("Introduction — L'annonce du sujet", 0, True, BLEU_FONCE),
    ("1.  Contexte et identification du problème", 0, False, GRIS),
    ("2.  Objectifs et résultats attendus", 0, False, GRIS),
    ("3.  Démarche méthodologique", 0, False, GRIS),
    ("4.  État des lieux des solutions existantes et solution proposée", 0, False, GRIS),
    ("5.  L'axe de différenciation (innovation technologique)", 0, False, GRIS),
    ("6.  Étude de faisabilité et conception de la solution", 0, False, GRIS),
]
droite = [
    ("7.  Fonctionnement de la solution : point de vue technique", 0, False, GRIS),
    ("8.  Démonstration", 0, True, ROUGE),
    ("9.  Marketing, vente et concurrence", 0, False, GRIS),
    ("10. Prévisions financières", 0, False, GRIS),
    ("11. Limites du projet et perspectives", 0, False, GRIS),
    ("", 0, False, GRIS),
    ("Conclusion", 0, True, BLEU_FONCE),
]
puces(s, gauche, x=Inches(0.6), y=Inches(1.6), w=Inches(6.3), taille=16, interligne=14)
puces(s, droite, x=Inches(7.0), y=Inches(1.6), w=Inches(6.0), taille=16, interligne=14)
notes(s, "Mon exposé suit le fil d'un projet d'entreprise : je pars du problème du "
         "marché, je présente la solution et sa technique, je la démontre, puis "
         "j'aborde son modèle économique avant de conclure sur les perspectives.")

# ====================================================================== #
#  INTRODUCTION                                                          #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, None, "Introduction — L'annonce du sujet")
puces(s, [
    ("Le sujet : concevoir, réaliser et déployer chez OPEN MOISE un système qui "
     "surveille en continu les réseaux supervisés, détecte automatiquement les "
     "pannes et alerte instantanément les équipes — 100 % open-source.", 0, False, GRIS),
    ("", 0, False, GRIS),
    ("Pourquoi ce thème ?", 0, True, BLEU_FONCE),
    ("Le réseau est le « système nerveux » des clients d'OPEN MOISE", 1, False, GRIS),
    ("Une minute d'indisponibilité menace les engagements de service (SLA)", 1, False, GRIS),
    ("Vécu de terrain chez OPEN MOISE : pannes des parcs clients découvertes "
     "trop tard, souvent signalées par le client lui-même", 1, False, GRIS),
    ("Conviction : l'open-source permet d'industrialiser la supervision sans licence", 1, False, GRIS),
], y=Inches(1.5), h=Inches(4.0), taille=19)
encadre(s, "Répondre à un besoin réel d'OPEN MOISE en alliant expertise réseau et "
           "démarche entrepreneuriale.", y=Inches(6.0))
notes(s, "J'ai choisi ce thème car il croise une réalité observée chez OPEN MOISE - "
         "les pannes des réseaux clients détectées trop tard - et une conviction : "
         "on peut industrialiser la supervision sans budget colossal grâce à "
         "l'open-source. C'est un besoin direct de mon entreprise d'accueil.")

# ====================================================================== #
#  PRÉSENTATION DE LA STRUCTURE D'ACCUEIL — OPEN MOISE                   #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, None, "Structure d'accueil — OPEN MOISE")
puces(s, [
    ("OPEN MOISE — Entreprise de Services du Numérique (ESN)", 0, True, BLEU_FONCE),
    ("", 0, False, GRIS),
    ("Activités : infogérance, intégration réseau & systèmes, développement, "
     "cybersécurité", 1, False, GRIS),
    ("Mission : garantir un SI performant, disponible et sécurisé à ses clients", 1, False, GRIS),
    ("[Effectif] collaborateurs · implantée à [ville] · depuis [année]", 1, False, GRIS),
    ("Projet mené au sein du [Pôle Infrastructures & Supervision]", 1, False, GRIS),
], y=Inches(1.5), h=Inches(3.6), taille=20)
encadre(s, "ESN qui supervise les réseaux de ses clients, OPEN MOISE est "
           "directement concernée : ce projet est un outil interne stratégique pour "
           "tenir ses engagements de service (SLA) et maîtriser ses coûts.",
        y=Inches(5.5), h=Inches(1.3))
notes(s, "Quelques mots sur mon entreprise d'accueil, OPEN MOISE : une ESN qui "
         "assure notamment l'infogérance et la supervision des réseaux de ses "
         "clients. Mon projet répond donc à un besoin direct de l'entreprise : "
         "détecter plus vite les pannes des parcs qu'elle exploite et respecter ses "
         "engagements de service.")

# ====================================================================== #
#  1 — CONTEXTE ET PROBLÈME                                             #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 1, "Contexte et identification du problème")
puces(s, [
    ("Contexte : OPEN MOISE, ESN d'infogérance, exploite et supervise les "
     "réseaux de plusieurs clients, avec des engagements de disponibilité "
     "contractuels (SLA, « 99,9 % »).", 0, False, GRIS),
    ("", 0, False, GRIS),
    ("Constat de terrain chez OPEN MOISE :", 0, True, BLEU_FONCE),
    ("Supervision largement manuelle et épisodique (ping à la main)", 1, False, GRIS),
    ("Pannes des parcs clients signalées par le client → réaction tardive", 1, False, GRIS),
    ("Aucune traçabilité des incidents → SLA difficiles à prouver", 1, False, GRIS),
    ("Solutions du marché jugées trop chères ou trop lourdes à déployer", 1, False, GRIS),
], y=Inches(1.5), h=Inches(4.8), taille=19)
notes(s, "Le contexte est celui d'OPEN MOISE : une ESN qui supervise les réseaux de "
         "ses clients sous engagement de service. Or cette supervision est encore "
         "largement manuelle. Résultat : on apprend la panne par le client lui-même, "
         "sans trace exploitable, ce qui met en péril le respect des SLA.")

# ---- PROBLÉMATIQUE (diapo dédiée) ----
s = ajouter_diapo(); fond(s, BLEU_FONCE)
tb = zone_texte(s, Inches(0.5), Inches(0.5), Inches(12.3), Inches(0.8))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "LA PROBLÉMATIQUE"
style_run(r, 24, True, MINT, police="Cambria")
boite = rectangle(s, Inches(1.0), Inches(1.7), Inches(11.3), Inches(2.6), BLANC)
tf = boite.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.4); tf.margin_right = Inches(0.4)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = ("« Comment OPEN MOISE peut-elle détecter en temps réel les "
                           "pannes des réseaux qu'elle supervise et réduire le temps "
                           "d'intervention de ses équipes, à l'aide d'une solution "
                           "automatisée, fiable et économiquement accessible ? »")
style_run(r, 21, True, NAVY, italique=True, police="Cambria")
tb = zone_texte(s, Inches(1.0), Inches(4.6), Inches(11.3), Inches(2.5))
tf = tb.text_frame
items = [
    ("Sous-questions :", True, MINT),
    ("Comment automatiser une surveillance continue et fiable des parcs clients ?", False, BLANC),
    ("Comment alerter les équipes d'astreinte d'OPEN MOISE, où qu'elles soient ?", False, BLANC),
    ("Comment offrir cette valeur sans alourdir les coûts d'OPEN MOISE ?", False, BLANC),
]
for i, (txt, g, c) in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(8)
    r = p.add_run(); r.text = ("" if i == 0 else "•  ") + txt
    style_run(r, 18, g, c)
notes(s, "Tout le projet répond à cette question centrale, formulée du point de vue "
         "d'OPEN MOISE : détecter en temps réel les pannes des réseaux supervisés et "
         "réduire le temps d'intervention des équipes, avec une solution automatisée, "
         "fiable et accessible. Les trois sous-questions guident ma conception.")

# ====================================================================== #
#  2 — OBJECTIFS ET RÉSULTATS                                           #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 2, "Objectifs et résultats attendus")
tableau(s, [
    ["Objectif", "Résultat attendu (indicateur)"],
    ["Détecter les pannes en temps réel", "MTTD < 2 min  (vs ~60 min)"],
    ["Alerter instantanément", "MTTR < 5 min  (vs ~45 min)"],
    ["Automatiser la surveillance", "0 intervention humaine sur la collecte"],
    ["Assurer la traçabilité", "100 % des événements journalisés"],
    ["Améliorer la disponibilité", "+2 pts  (~97,5 % → ~99,5 %)"],
    ["Maîtriser le coût", "0 FCFA de licence (open-source)"],
], y=Inches(1.5), h=Inches(3.8), taille=17, surligne_derniere_col=True)
encadre(s, "Objectif général : doter OPEN MOISE d'un outil fonctionnel, "
           "mesurablement plus performant que sa supervision manuelle actuelle, "
           "pour un coût négligeable.", y=Inches(5.7), h=Inches(1.1))
notes(s, "Mes objectifs sont chiffrés, donc vérifiables. Le cœur : faire chuter le "
         "temps de détection de l'ordre de l'heure à moins de deux minutes, et le "
         "temps de réaction à moins de cinq minutes, sans coût de licence.")

# ====================================================================== #
#  3 — DÉMARCHE MÉTHODOLOGIQUE                                          #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 3, "Démarche méthodologique")
puces(s, [
    ("Type : recherche appliquée / expérimentale — paradigme Design Science "
     "(on conçoit un artefact et on MESURE son efficacité).", 0, False, GRIS),
    ("", 0, False, GRIS),
    ("Cycle de travail itératif (inspiré DevOps) :", 0, True, BLEU_FONCE),
], y=Inches(1.5), h=Inches(1.8), taille=20)
boite = rectangle(s, Inches(0.9), Inches(3.5), Inches(11.5), Inches(0.9), GRIS_CLAIR)
tf = boite.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Analyse → Conception (UML/MERISE) → Développement → Test → Intégration → Mesure"
style_run(r, 17, True, BLEU_FONCE)
puces(s, [
    ("Outils & livrables : recherche documentaire · observation terrain · "
     "prototypage · collecte automatisée des métriques (journaux) · "
     "comparaison avant/après (manuel vs automatisé).", 0, False, GRIS),
], y=Inches(4.7), h=Inches(1.6), taille=19)
notes(s, "Ma démarche est celle de l'ingénieur-chercheur : je conçois un système "
         "réel et je prouve sa valeur par la mesure. J'ai travaillé en cycles "
         "itératifs, à la DevOps, et j'évalue par comparaison avant/après sur des "
         "indicateurs objectifs.")

# ====================================================================== #
#  4 — ÉTAT DES LIEUX                                                   #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 4, "État des lieux des solutions existantes")
tableau(s, [
    ["Solution", "Type", "Atout", "Frein majeur"],
    ["Nagios", "Open-source", "Éprouvé, extensible", "Config. complexe"],
    ["Zabbix", "Open-source", "Très complet", "Gourmand en ressources"],
    ["Centreon", "Mixte", "Interface soignée", "Modules avancés payants"],
    ["PRTG", "Propriétaire", "Simple", "Payant, Windows"],
    ["SolarWinds", "Propriétaire", "Haut de gamme", "Très coûteux"],
], y=Inches(1.5), h=Inches(3.4), taille=15)
encadre(s, "Le vide identifié : rien de léger + gratuit + simple + sur mesure, avec "
           "notification mobile native, adapté à OPEN MOISE et à ses clients PME.", y=Inches(5.4), h=Inches(1.1))
notes(s, "J'ai analysé les références du marché : soit des usines à gaz puissantes "
         "mais complexes, soit des produits simples mais propriétaires et chers. "
         "Pour OPEN MOISE et ses clients PME, il existe un vide pour une solution "
         "légère, gratuite et sur mesure.")

# ---- Solution proposée ----
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 4, "Solution proposée")
puces(s, [
    ("Un outil de supervision pour OPEN MOISE, automatisé, modulaire et "
     "100 % open-source, qui :", 0, True, BLEU_FONCE),
    ("découvre automatiquement les équipements des réseaux clients", 1, False, GRIS),
    ("sonde en continu la disponibilité (ICMP) et les services (TCP)", 1, False, GRIS),
    ("confirme la panne (anti-faux positif) et journalise tout (preuve de SLA)", 1, False, GRIS),
    ("alerte les équipes d'OPEN MOISE par e-mail ET Telegram (push mobile)", 1, False, GRIS),
    ("offre un tableau de bord web temps réel + rapports clients", 1, False, GRIS),
], y=Inches(1.5), h=Inches(3.6), taille=19)
encadre(s, "Stack : Ubuntu · Python · Cron · MySQL · Flask · Bootstrap · Bot Telegram",
        y=Inches(5.6), h=Inches(1.0), fond_c=BLEU_FONCE, barre=BLEU)
# corrige couleur texte de l'encadré (fond foncé)
notes(s, "Ma réponse au besoin d'OPEN MOISE : une solution complète mais légère. "
         "Elle découvre, surveille, confirme, journalise, alerte les équipes et "
         "affiche - le tout assemblé à partir de briques open-source éprouvées, et "
         "la journalisation sert aussi de preuve du respect des SLA.")

# ====================================================================== #
#  5 — DIFFÉRENCIATION                                                  #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 5, "L'axe de différenciation (innovation technologique)")
tableau(s, [
    ["Différenciateur", "Bénéfice pour OPEN MOISE"],
    ["Alerte Telegram native", "Astreinte alertée partout, sans coût de notification"],
    ["Ultra-légère", "1 VM par client / mutualisée : coûts d'exploitation réduits"],
    ["Code 100 % ouvert & maîtrisé", "Adaptable à chaque client, zéro vendor lock-in"],
    ["« Plug & supervise »", "Onboarding d'un nouveau client en < 30 min"],
    ["Prête pour l'IA", "Historique structuré → futur service prédictif vendable"],
], y=Inches(1.5), h=Inches(3.4), taille=15)
encadre(s, "Pour OPEN MOISE, l'innovation est l'ASSEMBLAGE : simplicité + gratuité + "
           "notification mobile + extensibilité IA → un avantage concurrentiel.", y=Inches(5.4), h=Inches(1.1))
notes(s, "Mon innovation ne réside pas dans un protocole nouveau, mais dans un "
         "positionnement unique : combiner simplicité, gratuité, alerte mobile "
         "instantanée et ouverture vers l'IA. Là où chaque concurrent impose un "
         "sacrifice, ma solution réunit ces atouts. Telegram natif est mon marqueur.")

# ====================================================================== #
#  6 — FAISABILITÉ & CONCEPTION                                        #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 6, "Étude de faisabilité et conception de la solution")
puces(s, [
    ("Faisabilité technique ✓ — briques matures, compétences déjà présentes "
     "chez OPEN MOISE, prototype réalisé et testé.", 0, False, VERT),
    ("Faisabilité économique ✓ — coût de licence nul, matériel mutualisé OPEN MOISE.", 0, False, VERT),
    ("Faisabilité organisationnelle ✓ — intégrable au processus d'astreinte d'OPEN MOISE.", 0, False, VERT),
    ("", 0, False, GRIS),
    ("Conception formalisée :", 0, True, BLEU_FONCE),
    ("UML : 7 diagrammes (contexte, cas d'usage, séquence, activité, classes, "
     "composants, déploiement)", 1, False, GRIS),
    ("MERISE : MCD → MLD → base MySQL (6 tables, dont une associative)", 1, False, GRIS),
    ("Architecture en 5 couches", 1, False, GRIS),
], y=Inches(1.5), h=Inches(5.2), taille=19)
notes(s, "Le projet est faisable sur les trois plans : technique - le prototype "
         "fonctionne -, économique - coût quasi nul - et organisationnel. J'ai "
         "rigoureusement conçu le système en UML et MERISE, sur une architecture en "
         "cinq couches.")

# ---- Architecture en 5 couches (IMAGE générée) ----
diapo_image(6, "Architecture en 5 couches",
            "architecture/architecture-5couches.png",
            legende="Découplage Moteur (Cron) ↔ Interface (Flask) : la surveillance "
                    "continue même si le tableau de bord est arrêté.",
            note_txt="Voici l'architecture en cinq couches. Le choix d'ingénierie le "
                     "plus important est le découplage entre la collecte et "
                     "l'affichage : ils communiquent uniquement par la base de "
                     "données, ce qui rend le système robuste.")

# ---- UML : 7 diagrammes (grille d'aperçu) ----
diapo_grille_images(6, "Conception UML — les 7 diagrammes", [
    ("architecture/diagrammes/diagramme-contexte.png", "1. Contexte"),
    ("architecture/diagrammes/cas-utilisation.png", "2. Cas d'utilisation"),
    ("architecture/diagrammes/diagramme-sequence.png", "3. Séquence"),
    ("architecture/diagrammes/diagramme-activite.png", "4. Activité"),
    ("architecture/diagrammes/diagramme-classes.png", "5. Classes"),
    ("architecture/diagrammes/diagramme-composants.png", "6. Composants"),
    ("architecture/diagrammes/diagramme-deploiement.png", "7. Déploiement"),
], note_txt="J'ai formalisé toute la conception avec sept diagrammes UML, du "
            "diagramme de contexte au diagramme de déploiement. Les versions "
            "pleine page figurent en annexe pour le détail.")

# ---- MERISE : MCD -> MLD ----
diapo_deux_images(6, "MERISE : du MCD au MLD (base MySQL, 6 tables)",
                  "database/mcd.png", "MCD — Modèle Conceptuel",
                  "database/mld.png", "MLD — Modèle Logique",
                  note_txt="Pour la base de données, j'ai suivi la méthode MERISE : "
                           "du modèle conceptuel - les entités et leurs associations, "
                           "dont la relation utilisateur-équipement - vers le modèle "
                           "logique - six tables, dont une table associative - "
                           "implémenté ensuite en MySQL.")

# ====================================================================== #
#  7 — FONCTIONNEMENT TECHNIQUE                                         #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 7, "Fonctionnement de la solution : point de vue technique")
etapes = [
    "1.  CRON déclenche le moteur toutes les 2 minutes",
    "2.  Charge les équipements (MySQL) puis sonde ICMP (ping) + TCP (service)",
    "3.  Met à jour le statut + journalise chaque vérification",
    "4.  Seuil d'échecs atteint ? (anti-faux positif)",
    "5.  ALERTE e-mail + Telegram  (avec anti-spam / dé-duplication)",
    "6.  Retour en ligne (UP) → notification de rétablissement automatique",
]
puces(s, [(e, 0, False, GRIS) for e in etapes], y=Inches(1.6), h=Inches(3.6), taille=19, interligne=12)
encadre(s, "Robustesse : seuil d'échecs (pas de fausse alerte), dé-duplication (pas "
           "de spam), sondes parallélisées (scan d'un /24 en ~12 s), journaux à rotation.",
        y=Inches(5.7), h=Inches(1.1))
notes(s, "Techniquement : toutes les deux minutes, Cron lance le moteur qui teste "
         "chaque équipement. Une panne n'est confirmée qu'après plusieurs échecs - "
         "pour éviter les fausses alertes - et une seule notification est envoyée "
         "par panne. Tout est parallélisé pour rester rapide.")

# ---- Sécurité by design ----
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 7, "Sécurité « by design »")
tableau(s, [
    ["Risque", "Parade implémentée"],
    ["Secrets dans le code", "Variables d'environnement (.env non versionné)"],
    ["Vol de mots de passe", "Hachage pbkdf2:sha256 (jamais en clair)"],
    ["Injection SQL", "Requêtes paramétrées systématiques"],
    ["Écoute des notifications", "SMTP STARTTLS + Telegram HTTPS"],
    ["Accès BDD trop large", "Compte applicatif au moindre privilège"],
    ["Accès non autorisé", "Authentification du tableau de bord"],
], y=Inches(1.5), h=Inches(4.0), taille=16)
notes(s, "La sécurité a été pensée dès la conception, pas ajoutée après : secrets "
         "hors du code, mots de passe hachés, requêtes paramétrées contre "
         "l'injection SQL, communications chiffrées et moindre privilège sur la base.")

# ====================================================================== #
#  8 — DÉMONSTRATION                                                    #
# ====================================================================== #
s = ajouter_diapo(); fond(s, NAVY)
cercle(s, Inches(11.2), Inches(-1.2), Inches(3.2), NAVY2)
rectangle(s, Inches(0.5), Inches(0.7), Inches(0.12), Inches(0.85), TEAL)
tb = zone_texte(s, Inches(0.8), Inches(0.65), Inches(12.0), Inches(0.95))
tf = tb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
r = p.add_run(); r.text = "8.  DÉMONSTRATION — scénario en direct"
style_run(r, 30, True, BLANC, police="Cambria")
etapes = [
    "Tableau de bord OPEN MOISE — parc client supervisé en temps réel (UP/KPI)",
    "Je simule une panne sur un équipement d'un réseau client",
    "Au cycle suivant → le statut bascule en DOWN (badge rouge)",
    "L'astreinte OPEN MOISE est alertée : e-mail + Telegram sur le téléphone",
    "Je rétablis → notification de rétablissement + acquittement auto",
    "Page Rapports : preuve de disponibilité / SLA à présenter au client",
]
tb = zone_texte(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.2))
tf = tb.text_frame
for i, e in enumerate(etapes):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(12)
    r = p.add_run(); r.text = f"{i+1}.  {e}"
    style_run(r, 20, False, BLANC)
encadre(s, "Plan B : captures d'écran + courte vidéo si la démo live est impossible.",
        y=Inches(6.4), h=Inches(0.8), fond_c=BLEU, barre=BLANC)
notes(s, "Pour la démonstration, je propose un scénario vivant tel qu'OPEN MOISE "
         "l'utiliserait : je montre le tableau de bord d'un parc client, je simule "
         "une panne, et vous verrez l'astreinte alertée par Telegram en temps réel. "
         "Puis je rétablis. J'ai un plan B en captures et vidéo.")

# ---- Démonstration : le TABLEAU DE BORD (capture hero) ----
diapo_image(8, "Démonstration — Tableau de bord temps réel",
            "soutenance/captures/capture-dashboard.png",
            legende="Vue d'ensemble OPEN MOISE : KPI, anneau de disponibilité (75 %), "
                    "alertes actives et état de chaque équipement supervisé.",
            note_txt="Voici le tableau de bord réel : en un coup d'œil, l'équipe "
                     "d'OPEN MOISE voit le nombre d'équipements en ligne, le taux de "
                     "disponibilité, les alertes actives et l'état détaillé du parc. "
                     "Ici, deux équipements sont en panne (DOWN), signalés en rouge.")

# ---- Démonstration : DÉTECTION + ALERTE (telegram + historique) ----
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 8, "Démonstration — détection et alerte instantanée")
# Notification Telegram (portrait) à gauche
_ch = os.path.join(RACINE, "soutenance/captures/capture-telegram.png")
w, h = _dim_ajustee(_ch, int(Inches(3.0)), int(Inches(5.2)))
s.shapes.add_picture(_ch, Inches(0.7), Inches(1.45), width=w, height=h)
tb = zone_texte(s, Inches(0.5), Inches(6.75), Inches(3.4), Inches(0.5))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Notification Telegram (astreinte OPEN MOISE)"
style_run(r, 12, True, BLEU_FONCE)
# Historique des alertes (paysage) à droite
_ca = os.path.join(RACINE, "soutenance/captures/capture-alertes.png")
w2, h2 = _dim_ajustee(_ca, int(Inches(8.3)), int(Inches(3.4)))
s.shapes.add_picture(_ca, Inches(4.4), Inches(1.7), width=w2, height=h2)
tb = zone_texte(s, Inches(4.4), Inches(1.45), Inches(8.3), Inches(0.35))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Historique des alertes (traçabilité + acquittement)"
style_run(r, 12, True, BLEU_FONCE)
encadre(s, "Chaîne complète : panne détectée → alerte e-mail + Telegram en < 2 min → "
           "acquittement → preuve de SLA dans l'historique.", x=Inches(4.4),
        y=Inches(5.5), w=Inches(8.3), h=Inches(1.2))
notes(s, "La détection déclenche en moins de deux minutes une alerte multi-canal. À "
         "gauche, la notification Telegram telle que la reçoit l'astreinte d'OPEN "
         "MOISE sur son smartphone. À droite, l'historique des alertes, qui assure la "
         "traçabilité et sert de preuve du respect des SLA auprès du client.")

# ---- Démonstration : page Équipements ----
diapo_image(8, "Démonstration — gestion des équipements",
            "soutenance/captures/capture-equipements.png",
            legende="Inventaire du parc client : découverte automatique, statut, "
                    "service supervisé et emplacement.",
            note_txt="La page de gestion des équipements présente l'inventaire du parc "
                     "client, alimenté automatiquement par la découverte réseau. "
                     "OPEN MOISE y ajoute ou retire des équipements en quelques clics.")

# ---- Démonstration : ce que chaque écran prouve (récap) ----
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 8, "Démonstration — ce que chaque écran prouve")
tableau(s, [
    ["Écran", "Ce qu'il prouve pour OPEN MOISE"],
    ["Tableau de bord", "Supervision temps réel, KPI, anneau de disponibilité"],
    ["Équipements", "Inventaire + découverte automatique des parcs clients"],
    ["Alertes", "Traçabilité + acquittement (preuve de SLA)"],
    ["Notification Telegram", "Astreinte alertée sur smartphone en < 2 min"],
    ["Rapports", "Statistiques & rapports de disponibilité exportables"],
], y=Inches(1.5), h=Inches(3.4), taille=16)
encadre(s, "Toutes ces captures proviennent de l'application réellement développée "
           "(Flask + Bootstrap), avec un parc client d'exemple.", y=Inches(5.4), h=Inches(1.0))
notes(s, "En résumé, chaque écran de l'application apporte une preuve concrète de "
         "valeur pour OPEN MOISE, de la supervision temps réel jusqu'à la preuve de "
         "SLA exportable pour le client.")

# ====================================================================== #
#  9 — MARKETING, VENTE & CONCURRENCE                                   #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 9, "Marketing, vente et concurrence")
puces(s, [
    ("Double valeur pour OPEN MOISE :", 0, True, BLEU_FONCE),
    ("Outil INTERNE : industrialise la supervision, fiabilise les SLA, réduit l'astreinte", 1, False, GRIS),
    ("Nouvelle OFFRE : « Supervision managée » vendue aux clients d'OPEN MOISE", 1, False, GRIS),
    ("", 0, False, GRIS),
    ("Proposition de valeur d'OPEN MOISE à ses clients :", 0, True, BLEU_FONCE),
    ("« OPEN MOISE surveille votre réseau 24/7 et intervient avant que vous ne "
     "constatiez la panne. »", 1, False, BLEU),
    ("", 0, False, GRIS),
    ("Monétisation (offre de service) :", 0, True, BLEU_FONCE),
    ("Abonnement mensuel de supervision managée (par site / par équipement)", 1, False, GRIS),
    ("Option Pro : SNMP, multi-sites, IA prédictive, rapports SLA avancés", 1, False, GRIS),
], y=Inches(1.45), h=Inches(5.2), taille=17)
notes(s, "Pour OPEN MOISE, la solution a une double valeur : en interne, elle "
         "industrialise la supervision et fiabilise les SLA ; en externe, elle "
         "devient une nouvelle offre de supervision managée facturée aux clients. "
         "C'est à la fois une économie et une source de revenus récurrents.")

# ---- Positionnement concurrentiel ----
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 9, "Positionnement concurrentiel")
tableau(s, [
    ["Critère", "Nagios/Zabbix", "PRTG/SolarWinds", "Solution OPEN MOISE"],
    ["Prix", "Gratuit mais coûteux", "Très coûteux", "Gratuit + services"],
    ["Simplicité", "Faible", "Élevée", "Élevée"],
    ["Légèreté", "Moyenne/Faible", "Faible", "Très élevée"],
    ["Alerte mobile native", "Plugin", "Appli", "Telegram natif"],
    ["Sur-mesure / ouvert", "Limité", "Non", "Total"],
], y=Inches(1.5), h=Inches(3.4), taille=15, surligne_derniere_col=True)
encadre(s, "Stratégie d'OPEN MOISE : ne pas affronter les géants sur le haut de gamme, "
           "mais se différencier sur le segment des PME avec une offre managée simple "
           "et sans licence (« océan bleu »).",
        y=Inches(5.4), h=Inches(1.1))
notes(s, "OPEN MOISE ne cherche pas à battre Zabbix sur les très grandes "
         "infrastructures. Sa stratégie est l'océan bleu : se différencier auprès de "
         "ses clients PME avec une offre de supervision managée simple et sans "
         "licence, là où les géants sont trop chers ou trop lourds.")

# ====================================================================== #
#  10 — PRÉVISIONS FINANCIÈRES                                          #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 10, "Prévisions financières (OPEN MOISE)")
tb = zone_texte(s, Inches(0.7), Inches(1.35), Inches(5.9), Inches(0.5))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "Investissement d'OPEN MOISE"
style_run(r, 17, True, BLEU_FONCE)
tableau(s, [
    ["Poste", "Coût"],
    ["Licences logicielles", "0 FCFA (open-source)"],
    ["Serveur / VM (mutualisé)", "~100 000 – 200 000 FCFA"],
    ["Développement & déploiement", "Interne (réalisé)"],
    ["Total d'entrée", "≈ 130 000 FCFA"],
], x=Inches(0.7), y=Inches(1.9), w=Inches(5.9), h=Inches(2.6), taille=15)
tb = zone_texte(s, Inches(6.9), Inches(1.35), Inches(5.9), Inches(0.5))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "Revenus & gains pour OPEN MOISE (an 1)"
style_run(r, 16, True, BLEU_FONCE)
tableau(s, [
    ["Source", "Hypothèse", "Montant"],
    ["Supervision managée", "10 clients × 25 000 FCFA/mois", "3 000 000 FCFA"],
    ["Mise en service client", "10 × 200 000 FCFA", "2 000 000 FCFA"],
    ["Option Pro (SNMP/IA)", "3 × 325 000 FCFA", "975 000 FCFA"],
    ["Pénalités SLA évitées", "estimation", "+ gains"],
    ["Total an 1", "", "≈ 6 000 000 FCFA"],
], x=Inches(6.9), y=Inches(1.9), w=Inches(5.9), h=Inches(2.9), taille=13, surligne_derniere_col=True)
encadre(s, "Double retour pour OPEN MOISE : revenus récurrents (offre managée) + "
           "économies internes (moins d'astreinte, pénalités SLA évitées). "
           "Investissement ≈ 130 000 FCFA amorti dès le 1er client.", y=Inches(5.85), h=Inches(0.95))
txt2 = zone_texte(s, Inches(0.7), Inches(6.95), Inches(12.0), Inches(0.35))
p = txt2.text_frame.paragraphs[0]
r = p.add_run(); r.text = ("Montants en francs CFA (XOF) — hypothèses indicatives à ajuster. "
                           "Parité fixe : 1 € = 655,957 FCFA.")
style_run(r, 10, False, GRIS, italique=True)
notes(s, "Pour OPEN MOISE, l'investissement est quasi nul : pas de licence, un "
         "serveur mutualisé, un développement déjà réalisé en interne. Le retour est "
         "double : des revenus récurrents via l'offre de supervision managée, et des "
         "économies internes - moins de temps d'astreinte et des pénalités SLA "
         "évitées. L'investissement est amorti dès le premier client.")

# ====================================================================== #
#  11 — LIMITES & PERSPECTIVES                                          #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, 11, "Limites du projet et perspectives")
puces(s, [
    ("Limites assumées :", 0, True, ROUGE),
    ("Supervision active → pas de métriques internes (CPU/RAM) sans SNMP", 1, False, GRIS),
    ("Réactivité bornée par l'intervalle de cycle (2 min)", 1, False, GRIS),
    ("Alertes dépendantes de la connectivité sortante", 1, False, GRIS),
    ("« Qui surveille le surveillant ? » → besoin de redondance", 1, False, GRIS),
    ("", 0, False, GRIS),
    ("Perspectives (feuille de route OPEN MOISE) :", 0, True, VERT),
    ("Intégration SNMP (performances détaillées des équipements clients)", 1, False, GRIS),
    ("Machine Learning : détection d'anomalies sur l'historique multi-clients", 1, False, GRIS),
    ("Maintenance prédictive vendable comme offre premium · Auto-remédiation (AIOps)", 1, False, GRIS),
    ("Portail multi-clients · Conteneurisation (Docker/K8s) · SMS, Slack, Teams", 1, False, GRIS),
], y=Inches(1.4), h=Inches(5.4), taille=18, interligne=6)
notes(s, "Je reste lucide sur les limites : pas encore de SNMP, réactivité bornée à "
         "deux minutes, et la question du superviseur lui-même. Mais ces limites "
         "tracent la feuille de route d'OPEN MOISE : SNMP, puis l'IA pour offrir à "
         "ses clients une supervision prédictive premium.")

# ====================================================================== #
#  CONCLUSION                                                           #
# ====================================================================== #
s = ajouter_diapo(); fond(s)
bandeau_titre(s, None, "Conclusion")
puces(s, [
    ("Pour OPEN MOISE : un outil fonctionnel, testé, 100 % open-source", 0, True, VERT),
    ("Détection temps réel (< 2 min) → SLA clients mieux respectés", 0, True, VERT),
    ("Double valeur : économies internes + nouvelle offre managée", 0, True, VERT),
    ("Une base extensible vers l'IA et la maintenance prédictive", 0, True, VERT),
], y=Inches(1.6), h=Inches(2.8), taille=21, interligne=14)
encadre(s, "Pour OPEN MOISE, l'automatisation transforme la supervision en avantage "
           "concurrentiel : de la réaction subie vers un service proactif, fiable et "
           "rentable pour ses clients.",
        y=Inches(5.0), h=Inches(1.6), fond_c=BLEU_FONCE, barre=BLEU)
notes(s, "En conclusion, ce projet dote OPEN MOISE d'un outil réel et performant, "
         "qui fiabilise ses engagements de service, génère des économies et ouvre "
         "une nouvelle offre commerciale. Pour l'entreprise, c'est un avantage "
         "concurrentiel concret et durable.")

# ---- Remerciements ----
s = ajouter_diapo(); fond(s, NAVY)
cercle(s, Inches(10.9), Inches(-1.3), Inches(3.6), NAVY2)
cercle(s, Inches(-1.0), Inches(5.4), Inches(3.2), NAVY2)
cercle(s, Inches(11.9), Inches(-0.4), Inches(1.5), TEAL_DK)
tb = zone_texte(s, Inches(1.0), Inches(2.8), Inches(11.3), Inches(2.0))
tf = tb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Merci de votre attention"
style_run(r, 40, True, BLANC, police="Cambria")
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Je me tiens à votre disposition pour vos questions"
style_run(r, 20, False, LIGHTTXT, italique=True)
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "\n[NOM Prénom] — Master 2 RIT — OPEN MOISE — 2026"
style_run(r, 16, True, MINT)
notes(s, "Je vous remercie de votre attention et je suis prêt à répondre à toutes "
         "vos questions, qu'elles soient techniques ou sur le volet économique.")

# ====================================================================== #
#  ANNEXES — diagrammes détaillés (diapos de secours pour le Q&A)        #
# ====================================================================== #
# Intercalaire
s = ajouter_diapo(); fond(s, NAVY)
cercle(s, Inches(11.0), Inches(-1.2), Inches(3.4), NAVY2)
cercle(s, Inches(-0.9), Inches(5.5), Inches(3.0), NAVY2)
rectangle(s, Inches(5.4), Inches(4.35), Inches(2.6), Inches(0.06), TEAL)
tb = zone_texte(s, Inches(1.0), Inches(3.0), Inches(11.3), Inches(1.5))
tf = tb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "ANNEXES — Diagrammes de conception détaillés"
style_run(r, 30, True, BLANC, police="Cambria")
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Diapositives de secours pour les questions du jury"
style_run(r, 16, False, LIGHTTXT, italique=True)
notes(s, "Voici en annexe les diagrammes détaillés, que je peux afficher pour "
         "répondre précisément à vos questions sur la conception.")

# UML pleine page (1 par diapo)
_uml = [
    ("architecture/diagrammes/diagramme-contexte.png", "Annexe — Diagramme de contexte"),
    ("architecture/diagrammes/cas-utilisation.png", "Annexe — Diagramme de cas d'utilisation"),
    ("architecture/diagrammes/diagramme-sequence.png", "Annexe — Diagramme de séquence"),
    ("architecture/diagrammes/diagramme-activite.png", "Annexe — Diagramme d'activité"),
    ("architecture/diagrammes/diagramme-classes.png", "Annexe — Diagramme de classes"),
    ("architecture/diagrammes/diagramme-composants.png", "Annexe — Diagramme de composants"),
    ("architecture/diagrammes/diagramme-deploiement.png", "Annexe — Diagramme de déploiement"),
]
for chemin, titre in _uml:
    diapo_image(None, titre, chemin)

# MLD pleine page (schéma relationnel détaillé)
diapo_image(None, "Annexe — Modèle Logique de Données (MLD)", "database/mld.png",
            legende="6 tables : utilisateurs, equipements, journaux, alertes, "
                    "responsabilite (associative), statistiques — avec clés primaires "
                    "(PK) et étrangères (FK).")

# ---------------------------------------------------------------------- #
sortie = os.path.join(os.path.dirname(__file__), "presentation-projet.pptx")
prs.save(sortie)
print(f"PowerPoint généré : {sortie}  ({len(prs.slides.__iter__.__self__._sldIdLst)} diapositives)")
